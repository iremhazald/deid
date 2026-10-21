from __future__ import annotations

import html
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document

from app.core.config import get_settings
from app.detectors.registry import PageText
from app.services.docx_preflight import UploadRejected, preflight_docx_path
from app.services.jobs import SpanBox

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _w(local: str) -> str:
    return f"{{{_W}}}{local}"


_SKIP_TAGS = {
    _w("drawing"),
    _w("pict"),
    _w("object"),
    _w("txbxContent"),
    f"{{{_M}}}oMath",
    f"{{{_M}}}oMathPara",
}
_NOTICE_TAGS = {
    _w("commentReference"),
    _w("footnoteReference"),
    _w("endnoteReference"),
}
_UNSUPPORTED_PARTS = (
    "word/media/",
    "word/embeddings/",
    "word/diagrams/",
    "word/comments.xml",
    "word/footnotes.xml",
    "word/endnotes.xml",
)

_ocr_engine: Any | None = None


def _get_ocr_engine() -> Any:
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr import RapidOCR

        _ocr_engine = RapidOCR()
    return _ocr_engine


def extract_pdf_with_ocr(
    document: pymupdf.Document,
) -> tuple[list[PageText], str, list[SpanBox], bool]:
    pages: list[PageText] = []
    chunks: list[str] = []
    span_boxes: list[SpanBox] = []
    document_offset = 0
    ocr_used = False

    for page_number in range(document.page_count):
        page = document[page_number]
        native_text = str(page.get_text("text"))
        if native_text.strip():
            page_text = native_text
        else:
            ocr_used = True
            zoom = 2.0
            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(zoom, zoom),
                colorspace=pymupdf.csRGB,
                alpha=False,
            )
            result = _get_ocr_engine()(pixmap.tobytes("png"))
            texts = list(result.txts) if result.txts is not None else []
            boxes = list(result.boxes) if result.boxes is not None else []
            line_chunks: list[str] = []
            local_offset = 0
            for text, box in zip(texts, boxes, strict=False):
                clean = str(text).strip()
                if not clean:
                    continue
                line_chunks.append(clean)
                points = list(box)
                x_values = [float(point[0]) / zoom for point in points]
                y_values = [float(point[1]) / zoom for point in points]
                span_boxes.append(
                    SpanBox(
                        page=page_number,
                        start=document_offset + local_offset,
                        end=document_offset + local_offset + len(clean),
                        rectangle=(
                            min(x_values),
                            min(y_values),
                            max(x_values),
                            max(y_values),
                        ),
                    )
                )
                local_offset += len(clean) + 1
            page_text = "\n".join(line_chunks)
            if page_text:
                page_text += "\n"

        pages.append(PageText(page_number, page_text, document_offset))
        chunks.append(page_text)
        document_offset += len(page_text)

    return pages, "".join(chunks), span_boxes, ocr_used


def pdf_text_layouts(document: pymupdf.Document, pages: list[PageText]) -> list[dict[str, Any]]:
    """Presentation boxes for the existing page text. Does not change that text."""
    layouts: list[dict[str, Any]] = []
    for item in pages:
        if item.page < 0 or item.page >= document.page_count:
            continue
        try:
            layout = align_pdf_page_layout(document[item.page], item.text)
        except (RuntimeError, ValueError, TypeError):
            layout = None
        if layout is None:
            continue
        layouts.append({"page": item.page, **layout})
    return layouts


def align_pdf_page_layout(page: Any, page_text: str) -> dict[str, Any] | None:
    """Map PDF span boxes onto ``page_text`` without rewriting it.

    ``get_text("text")`` turns side-by-side spans into separate lines. The boxes
    still share a baseline, so the text workspace can place those same characters
    back on that line. A span that cannot be found in order abandons the page.
    """
    if not page_text:
        return None
    try:
        info = page.get_text("dict") or {}
    except (RuntimeError, ValueError, TypeError):
        return None
    runs: list[dict[str, Any]] = []
    cursor = 0
    for block in info.get("blocks") or []:
        if block.get("type") != 0:
            continue
        for line in block.get("lines") or []:
            for span in line.get("spans") or []:
                raw = str(span.get("text") or "")
                if raw == "":
                    continue
                index = page_text.find(raw, cursor)
                if index < 0:
                    return None
                bbox = span.get("bbox") or (0, 0, 0, 0)
                if len(bbox) < 4:
                    return None
                x0, y0, x1, y1 = (float(bbox[0]), float(bbox[1]), float(bbox[2]), float(bbox[3]))
                if x1 < x0 or y1 < y0:
                    return None
                flags = int(span.get("flags") or 0)
                size = float(span.get("size") or 0) or max(y1 - y0, 1.0)
                runs.append(
                    {
                        "start": index,
                        "end": index + len(raw),
                        "x0": round(x0, 2),
                        "y0": round(y0, 2),
                        "x1": round(x1, 2),
                        "y1": round(y1, 2),
                        "size": round(size, 2),
                        "bold": bool(flags & 16),
                    }
                )
                cursor = index + len(raw)
    if not runs:
        return None
    rect = page.rect
    width = float(rect.width)
    height = float(rect.height)
    if width <= 0 or height <= 0:
        return None
    return {"width": round(width, 2), "height": round(height, 2), "runs": runs}


@dataclass(frozen=True)
class DocxText:
    text: str
    warnings: tuple[str, ...]


def extract_docx(path: Path) -> str:
    return extract_docx_text(path).text


def extract_docx_text(path: Path, *, max_chars: int | None = None) -> DocxText:
    """Read a DOCX as normalized text in body order.

    Headers and footers are appended after the body. Their exact position on
    the page is not reconstructed. Images, text boxes, comments, footnotes,
    equations, SmartArt, and embedded objects are not extracted.
    """
    preflight_docx_path(path)
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ValueError("DOCX_CORRUPT") from exc

    flags: set[str] = set()
    if _package_has_unsupported_parts(path):
        flags.add("unsupported")
    blocks = _blocks(document.element.body, flags)
    headers, footers = _section_notes(document, flags)
    parts = ["\n".join(piece) for piece in (blocks, headers, footers) if piece]
    text = "\n".join(parts)
    limit = max_chars if max_chars is not None else get_settings().redaction_max_extracted_chars
    if len(text) > limit:
        raise UploadRejected(413, "DOCX_LIMIT_EXCEEDED")
    warnings = ("DOCX_UNSUPPORTED_CONTENT",) if "unsupported" in flags else ()
    return DocxText(text, warnings)


def _package_has_unsupported_parts(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            name = info.filename.replace("\\", "/").lower()
            if any(name == part or name.startswith(part) for part in _UNSUPPORTED_PARTS):
                return True
    return False


def _section_notes(document: Any, flags: set[str]) -> tuple[list[str], list[str]]:
    headers: list[str] = []
    footers: list[str] = []
    for section in document.sections:
        headers.extend(_story_blocks(section.header, flags))
        footers.extend(_story_blocks(section.footer, flags))
        if section.different_first_page_header_footer:
            headers.extend(_story_blocks(section.first_page_header, flags))
            footers.extend(_story_blocks(section.first_page_footer, flags))
    return headers, footers


def _story_blocks(part: Any, flags: set[str]) -> list[str]:
    if part is None or bool(getattr(part, "is_linked_to_previous", False)):
        return []
    element = getattr(part, "_element", None)
    if element is None:
        return []
    return _blocks(element, flags)


def _blocks(parent: Any, flags: set[str]) -> list[str]:
    blocks: list[str] = []
    for child in parent:
        if child.tag == _w("p"):
            text = _paragraph_text(child, flags)
            if _is_list_paragraph(child) and text and not text.startswith(("- ", "• ")):
                text = f"- {text}"
            if text:
                blocks.append(text)
        elif child.tag == _w("tbl"):
            lines = _table_lines(child, flags)
            if any(line.strip() for line in lines):
                blocks.append("\n".join(lines))
        elif child.tag == _w("sdt"):
            content = child.find(_w("sdtContent"))
            if content is not None:
                blocks.extend(_blocks(content, flags))
        elif child.tag in _SKIP_TAGS:
            flags.add("unsupported")
        elif any(node.tag in _SKIP_TAGS or node.tag in _NOTICE_TAGS for node in child.iter()):
            flags.add("unsupported")
    return blocks


def _is_list_paragraph(paragraph: Any) -> bool:
    properties = paragraph.find(_w("pPr"))
    return properties is not None and properties.find(_w("numPr")) is not None


def _paragraph_text(paragraph: Any, flags: set[str]) -> str:
    parts: list[str] = []

    def walk(node: Any, in_deletion: bool) -> None:
        for child in node:
            if child.tag in _SKIP_TAGS:
                flags.add("unsupported")
                continue
            if child.tag == _w("instrText"):
                continue
            if child.tag in _NOTICE_TAGS:
                flags.add("unsupported")
                continue
            if child.tag == _w("del"):
                walk(child, True)
                continue
            if not in_deletion:
                if child.tag == _w("t") and child.text:
                    parts.append(child.text)
                elif child.tag == _w("tab"):
                    parts.append("\t")
                elif child.tag in {_w("br"), _w("cr")}:
                    parts.append("\n")
            walk(child, in_deletion)

    walk(paragraph, False)
    return "".join(parts).strip()


def _table_lines(table: Any, flags: set[str]) -> list[str]:
    lines: list[str] = []
    for row in table.findall(_w("tr")):
        cells = [_cell_text(cell, flags) for cell in row.findall(_w("tc"))]
        lines.append("\t".join(cells))
    return lines


def _cell_text(cell: Any, flags: set[str]) -> str:
    bits: list[str] = []
    for child in cell:
        if child.tag == _w("p"):
            text = _paragraph_text(child, flags)
            if text:
                bits.append(text)
        elif child.tag == _w("tbl"):
            nested = _table_lines(child, flags)
            if nested:
                bits.append(" ".join(line.replace("\t", " ") for line in nested))
    return " ".join(bits)


def redact_text_content(
    text: str,
    spans: list[tuple[int, int, str]],
    manual_values: list[str],
    method: str,
) -> tuple[str, list[str]]:
    selected_values: list[str] = []
    replacements: list[tuple[int, int, str]] = []
    counters: dict[str, int] = {}

    for start, end, entity_type in spans:
        original = text[start:end]
        selected_values.append(original)
        counters[entity_type] = counters.get(entity_type, 0) + 1
        replacement = (
            f"[{entity_type}_{counters[entity_type]}]"
            if method == "LABEL"
            else "█" * max(1, len(original))
        )
        replacements.append((start, end, replacement))

    for value in manual_values:
        cursor = 0
        found_any = False
        while True:
            start = text.find(value, cursor)
            if start < 0:
                break
            found_any = True
            selected_values.append(value)
            replacement = (
                "[MANUAL_REDACTION]"
                if method == "LABEL"
                else "█" * max(1, len(value))
            )
            replacements.append((start, start + len(value), replacement))
            cursor = start + len(value)
        if not found_any:
            raise ValueError("SPAN_MAPPING_FAILED")

    # Remove overlapping operations conservatively, then edit back-to-front so
    # offsets remain stable.
    accepted: list[tuple[int, int, str]] = []
    for operation in sorted(replacements, key=lambda item: (item[0], -item[1])):
        if any(operation[0] < other[1] and operation[1] > other[0] for other in accepted):
            continue
        accepted.append(operation)
    output = text
    for start, end, replacement in sorted(accepted, reverse=True):
        output = output[:start] + replacement + output[end:]
    return output, selected_values


def write_text_output(text: str, path: Path, output_format: str) -> None:
    if output_format == "TXT":
        path.write_text(text, encoding="utf-8")
        return
    if output_format == "DOCX":
        document = Document()
        for line in text.splitlines() or [""]:
            document.add_paragraph(line)
        document.core_properties.title = "Redacted legal document"
        document.core_properties.author = ""
        document.core_properties.subject = ""
        document.core_properties.comments = ""
        document.save(str(path))
        return
    if output_format != "PDF":
        raise ValueError("OUTPUT_FORMAT_INVALID")

    document = pymupdf.open()
    lines = text.splitlines() or [""]
    for start in range(0, len(lines), 45):
        page = document.new_page()
        body = html.escape("\n".join(lines[start : start + 45]))
        page.insert_htmlbox(
            pymupdf.Rect(50, 50, page.rect.width - 50, page.rect.height - 50),
            f"<pre style='font-family:sans-serif;font-size:10pt'>{body}</pre>",
        )
    document.save(path, garbage=4, deflate=True, clean=True)
    document.close()


def extract_generated_output(path: Path, output_format: str) -> str:
    if output_format == "TXT":
        return path.read_text(encoding="utf-8")
    if output_format == "DOCX":
        return extract_docx(path)
    document = pymupdf.open(path)
    try:
        return "".join(
            str(document[index].get_text("text"))
            for index in range(document.page_count)
        )
    finally:
        document.close()
